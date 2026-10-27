"""
convert_excel.py - OPSIONAL
Membaca file Excel Monthly Report lalu menulis ulang js/data.js.

Cara pakai (di terminal VS Code, dari folder proyek):
    pip install openpyxl
    python tools/convert_excel.py "path/ke/file.xlsx"

Aturan:
- Bulan ditentukan dari kolom BLTH (kata pertama, mis. "JUNI 2026" -> bulan 6).
- Hanya angka yang ada di Excel yang ditulis. Sel kosong / "-" menjadi null.
- Sheet yang tidak berisi bulan tertentu tidak ditulis untuk bulan itu.
"""
import json, sys, re
from pathlib import Path
import openpyxl

BULAN = ["JANUARI","FEBRUARI","MARET","APRIL","MEI","JUNI","JULI","AGUSTUS",
         "SEPTEMBER","OKTOBER","NOVEMBER","DESEMBER"]
NAMA = [b.capitalize() for b in BULAN]


def num(v):
    """Ubah nilai sel menjadi angka; '-' / kosong -> None."""
    if v is None or (isinstance(v, str) and v.strip() in ("", "-")):
        return None
    if isinstance(v, str):
        try:
            return float(v.replace(".", "").replace(",", "."))
        except ValueError:
            return None
    return int(v) if float(v).is_integer() else float(v)


def month_of(blth):
    if not isinstance(blth, str):
        return None
    w = blth.strip().split()[0].upper()
    return BULAN.index(w) + 1 if w in BULAN else None


def rows(ws, col_map, min_row=2):
    """Baca baris per bulan. col_map: {nama_field: (index_kolom_0, tipe)}"""
    out = {}
    for r in ws.iter_rows(min_row=min_row, values_only=True):
        m = month_of(r[1]) if len(r) > 1 else None
        if not m:
            continue
        item = {}
        for key, (idx, kind) in col_map.items():
            v = r[idx] if idx < len(r) else None
            if kind == "n":
                item[key] = num(v)
            elif isinstance(v, float) and v.is_integer():
                item[key] = str(int(v))
            else:
                item[key] = str(v).strip() if v is not None else None
        out.setdefault(m, []).append(item)
    return out


def build(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    data = {}

    def put(m, key, val):
        data.setdefault(m, {})[key] = val

    es_cols = {"nama": (2, "s"), "sangatTidakPuas": (3, "n"), "tidakPuas": (4, "n"),
               "kurangPuas": (5, "n"), "cukupPuas": (6, "n"), "puas": (7, "n"),
               "sangatPuas": (8, "n"), "skala5": (9, "n"), "skala100": (10, "n"),
               "status": (11, "s")}
    if "ESURVEY" in wb.sheetnames:
        for m, v in rows(wb["ESURVEY"], es_cols).items():
            put(m, "esurveyKanwil", v)
    if "ESURVEY CABANG" in wb.sheetnames:
        for m, v in rows(wb["ESURVEY CABANG"], es_cols).items():
            put(m, "esurveyCabang", v)

    if "DIGITAL SERVICE ALL KANWIL" in wb.sheetnames:
        c = {"peringkat": (0, "n"), "kode": (2, "s"), "nama": (3, "s"), "bobot": (6, "n"),
             "target": (7, "n"), "realisasi": (8, "n"), "pencapaian": (9, "n")}
        for m, v in rows(wb["DIGITAL SERVICE ALL KANWIL"], c).items():
            put(m, "digitalKanwil", v)

    if "DIGITAL SERVICE CABANG" in wb.sheetnames:
        c = {"kode": (2, "s"), "nama": (3, "s"),
             "targetJHT": (4, "n"), "realJHT": (5, "n"), "pctRealJHT": (6, "n"), "pctCapJHT": (7, "n"),
             "targetJP": (8, "n"), "realJP": (9, "n"), "pctRealJP": (10, "n"), "pctCapJP": (11, "n"),
             "targetBea": (12, "n"), "realBea": (13, "n"), "pctRealBea": (14, "n"), "pctCapBea": (15, "n"),
             "hasil": (16, "n")}
        for m, v in rows(wb["DIGITAL SERVICE CABANG"], c).items():
            put(m, "digitalCabang", v)

    if "SLA KLAIM" in wb.sheetnames:
        c = {"kode": (2, "s"), "nama": (3, "s"), "jht": (4, "n"), "jkk": (5, "n"), "jkm": (6, "n"),
             "jpn": (7, "n"), "total": (8, "n"), "slaJht": (9, "n"), "slaJkk": (10, "n"),
             "slaJkm": (11, "n"), "slaJpn": (12, "n"), "totalSla": (13, "n"),
             "pencapaian": (14, "n"), "bobot": (15, "n"), "skor": (16, "n")}
        for m, v in rows(wb["SLA KLAIM"], c).items():
            put(m, "slaKlaim", v)

    if "PENGADUAN" in wb.sheetnames:
        c = {"kode": (2, "s"), "nama": (3, "s"), "total": (4, "n"), "sesuaiSla": (5, "n"),
             "pencapaian": (6, "n"), "bobot": (7, "n"), "skor": (8, "n")}
        res = rows(wb["PENGADUAN"], c)
        for m, v in res.items():
            for it in v:
                # Excel menyimpan 3.142 untuk 3142 (titik ribuan terbaca desimal)
                for k in ("total", "sesuaiSla"):
                    if isinstance(it[k], float) and not float(it[k]).is_integer():
                        it[k] = int(round(it[k] * 1000))
            put(m, "pengaduan", v)

    if "MLT" in wb.sheetnames:
        c = {"kode": (3, "s"), "nama": (2, "s"), "targetUnit": (4, "n"), "targetNominal": (5, "n"),
             "targetUnit120": (6, "n"), "targetNominal120": (7, "n"), "realUnit": (8, "n"),
             "realNominal": (9, "n"), "pctUnit": (10, "n"), "pctNominal": (11, "n")}
        for m, v in rows(wb["MLT"], c).items():
            put(m, "mlt", v)
    return data


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    data = build(sys.argv[1])
    target = Path(__file__).resolve().parent.parent / "js" / "data.js"
    old = target.read_text(encoding="utf-8") if target.exists() else ""
    head = old.split("const REPORT_DATA")[0] if "const REPORT_DATA" in old else ""
    # pertahankan 'pdf' dan 'catatan' yang sudah Anda isi manual
    prev = {}
    if "const REPORT_DATA" in old:
        try:
            prev = json.loads(old.split("const REPORT_DATA = ", 1)[1].rsplit(";", 1)[0])
        except Exception:
            prev = {}
    out = {}
    for m in range(1, 13):
        if m in data:
            p = prev.get(str(m), {})
            out[str(m)] = {"periode": f"{NAMA[m-1]} 2026", "pdf": p.get("pdf"),
                           "catatan": p.get("catatan", []), "sections": data[m]}
    target.write_text(head + "const REPORT_DATA = " + json.dumps(out, ensure_ascii=False, indent=1) + ";\n",
                      encoding="utf-8")
    print("Bulan yang ditulis:", [NAMA[int(k)-1] for k in out])


if __name__ == "__main__":
    main()
